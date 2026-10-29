"""F33: Chinese journal manuscripts as Word users write them — heading
styles, superscript [n] citations, the full-width full stop "．"."""

from docx import Document

from citecheck.parse.parser import parse_docx
from citecheck.support.claims import split_sentences


def _doc(tmp_path, build):
    doc = Document()
    build(doc)
    path = tmp_path / "manuscript.docx"
    doc.save(path)
    return parse_docx(path)


def _cited(p, text, label):
    """Paragraph ``text`` followed by a superscript citation ``label``."""
    p.add_run(text)
    p.add_run(label).font.superscript = True


def _refs(doc, n):
    doc.add_paragraph("参考文献", style="Heading 1")
    for i in range(1, n + 1):
        doc.add_paragraph(f"[{i}] 张三, 李四. 题名{i}[J]. 计算机学报, 2020, 43(1): 1-10.")


def test_styled_headings_turn_off_numbered_text_headings(tmp_path):
    def build(doc):
        doc.add_paragraph("知识图谱构建技术综述", style="Title")
        doc.add_paragraph("1 知识图谱的定义与架构", style="Heading 1")
        doc.add_paragraph("知识图谱是结构化的语义知识库．")
        doc.add_paragraph("2 知识图谱的构建技术", style="Heading 1")
        doc.add_paragraph("2.1 信息抽取", style="Heading 2")
        doc.add_paragraph("1）实体消歧")  # a numbered list item, not a section
        doc.add_paragraph("实体消歧用于解决同名实体产生歧义的问题．")
        _refs(doc, 1)

    parsed = _doc(tmp_path, build)
    titles = [s.title for s in parsed.sections]
    assert titles == ["", "1 知识图谱的定义与架构", "2 知识图谱的构建技术", "参考文献"]


def test_unstyled_document_still_reads_numbered_headings(tmp_path):
    def build(doc):
        doc.add_paragraph("任务型对话系统研究综述")
        doc.add_paragraph("1 引言")
        doc.add_paragraph("对话系统受到广泛关注．")
        doc.add_paragraph("2 对话系统类型")
        doc.add_paragraph("对话系统分为两类．")
        _refs(doc, 1)

    parsed = _doc(tmp_path, build)
    assert [s.title for s in parsed.sections][1:3] == ["1 引言", "2 对话系统类型"]


def test_fullwidth_full_stop_ends_a_sentence():
    sents = [s.text for s in split_sentences("本文综述了目标检测方法．实验表明该方法有效．")]
    assert sents == ["本文综述了目标检测方法．", "实验表明该方法有效．"]


def test_superscripts_before_first_heading_kept_in_short_prose(tmp_path):
    """An unheaded introduction written with "．" is prose, not an author
    line: its superscript citations count."""
    def build(doc):
        doc.add_paragraph("知识图谱构建技术综述", style="Title")
        p = doc.add_paragraph()
        p.add_run("刘峤")
        p.add_run("1,2").font.superscript = True  # affiliation, not a citation
        _cited(doc.add_paragraph(), "语义网络技术向用户提供查询环境", "[1]")
        doc.paragraphs[-1].add_run("．其核心是以图形方式返回知识．")
        doc.add_paragraph("1 知识图谱的定义与架构", style="Heading 1")
        _cited(doc.add_paragraph(), "知识图谱的概念由谷歌提出", "[2]")
        doc.paragraphs[-1].add_run("．")
        _refs(doc, 2)

    parsed = _doc(tmp_path, build)
    assert [m.raw for m in parsed.markers] == ["[1]", "[2]"]
    assert [m.ref_ids for m in parsed.markers] == [["r1"], ["r2"]]


def test_numbered_heading_ends_front_matter(tmp_path):
    """No canonical 引言 heading: the first numbered top-level heading
    starts the body, so a short line under it keeps its citation."""
    def build(doc):
        doc.add_paragraph("知识图谱构建技术综述", style="Title")
        doc.add_paragraph("1 知识图谱的定义与架构", style="Heading 1")
        _cited(doc.add_paragraph(), "知识图谱的三层架构", "[1]")
        doc.add_paragraph("2 结束语", style="Heading 1")  # the only canonical one
        doc.add_paragraph("本文总结了知识图谱构建技术．")
        _refs(doc, 1)

    parsed = _doc(tmp_path, build)
    assert [m.raw for m in parsed.markers] == ["[1]"]
