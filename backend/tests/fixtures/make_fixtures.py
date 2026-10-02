"""Deterministic generator for the test fixtures.

Run:  python tests/fixtures/make_fixtures.py
Writes <name>.docx and <name>.expected.json side by side. References are
real papers (they get verified against Crossref in T2). Marker text is
deliberately split across runs with mixed bold/italic to mimic Word.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document
from docx.table import Table
from lxml import etree

HERE = Path(__file__).parent
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
ns = f'xmlns:w="{W}"'


def qn(tag: str) -> str:
    return f"{{{W}}}{tag}"


# ---------------------------------------------------------------- helpers


def add_para(target, parts, style: str | None = None):
    """parts: list of (text, fmt) with fmt keys b/i/sup; '\t' -> w:tab."""
    p = target.add_paragraph(style=style)
    for text, fmt in parts:
        r = p.add_run(text)
        if fmt.get("b"):
            r.bold = True
        if fmt.get("i"):
            r.italic = True
        if fmt.get("sup"):
            r.font.superscript = True
    return p


def add_heading(target, text: str, level: int = 1):
    return target.add_paragraph(text, style=f"Heading {level}")


def set_pstyle(p, style_id: str):
    ppr = p._p.get_or_add_pPr()
    st = etree.SubElement(ppr, qn("pStyle"))
    st.set(qn("val"), style_id)


def inject_localized_styles(doc):
    """styleId '1'/'2' whose w:name is 'heading 1'/'heading 2' — how zh
    Word names built-in heading styles."""
    styles_el = doc.styles.element
    for sid, lvl in (("1", 0), ("2", 1)):
        el = etree.fromstring(
            f'<w:style {ns} w:type="paragraph" w:styleId="{sid}">'
            f'<w:name w:val="heading {lvl + 1}"/>'
            f'<w:pPr><w:outlineLvl w:val="{lvl}"/></w:pPr></w:style>'
        )
        styles_el.append(el)


def zotero_instr(item_data: list[dict], formatted: str) -> str:
    payload = {
        "citationID": "xK9fT2mQ",
        "properties": {"formattedCitation": formatted, "plainCitation": formatted},
        "citationItems": [
            {"id": 100 + i, "itemData": d} for i, d in enumerate(item_data)
        ],
        "schema": "https://github.com/citation-style-language/schema/raw/master/csl-citation.json",
    }
    return " ZOTERO_ITEM CSL_CITATION " + json.dumps(payload, ensure_ascii=False) + " "


def add_complex_field(p, instr: str, visible: str):
    """Append a complex field (begin/instr/separate/result/end) to a paragraph."""
    for frag in (
        '<w:fldChar w:fldCharType="begin"/>',
        f'<w:instrText xml:space="preserve">{_esc(instr)}</w:instrText>',
        '<w:fldChar w:fldCharType="separate"/>',
        f"<w:t>{_esc(visible)}</w:t>",
        '<w:fldChar w:fldCharType="end"/>',
    ):
        r = etree.fromstring(f"<w:r {ns}>{frag}</w:r>")
        p._p.append(r)


def _esc(s: str) -> str:
    return (
        s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )


def zotero_item(title, authors, year, doi=None, venue=None, item_type="article-journal"):
    d = {
        "id": 1,
        "type": item_type,
        "title": title,
        "author": [{"family": f, "given": g} for f, g in authors],
        "issued": {"date-parts": [[year]]},
    }
    if doi:
        d["DOI"] = doi
    if venue:
        d["container-title"] = venue
    return d


# ---------------------------------------------------------------- references

EN_REFS = [
    "[1] Vaswani A, Shazeer N, Parmar N, et al. Attention is all you need. "
    "Advances in Neural Information Processing Systems. 2017;30:5998-6008.",
    "[2] Devlin J, Chang M-W, Lee K, Toutanova K. BERT: Pre-training of deep "
    "bidirectional transformers for language understanding. Proceedings of "
    "NAACL-HLT. 2019:4171-4186. doi:10.18653/v1/N19-1423.",
    "[3] Brown TB, Mann B, Ryder N, et al. Language models are few-shot "
    "learners. Advances in Neural Information Processing Systems. "
    "2020;33:1877-1901.",
    "[4] Lewis P, Perez E, Piktus A, et al. Retrieval-augmented generation for "
    "knowledge-intensive NLP tasks. Advances in Neural Information Processing "
    "Systems. 2020;33:9459-9474.",
    "[5] Ji Z, Lee N, Frieske R, et al. Survey of hallucination in natural "
    "language generation. ACM Computing Surveys. 2023;55(12):1-38. "
    "doi:10.1145/3571730.",
    "[6] He K, Zhang X, Ren S, Sun J. Deep residual learning for image "
    "recognition. Proceedings of the IEEE Conference on Computer Vision and "
    "Pattern Recognition. 2016:770-778. doi:10.1109/CVPR.2016.90.",
    "[7] Mikolov T, Chen K, Corrado G, Dean J. Efficient estimation of word "
    "representations in vector space. arXiv preprint arXiv:1301.3781. 2013.",
    "[8] Sutskever I, Vinyals O, Le QV. Sequence to sequence learning with "
    "neural networks. Advances in Neural Information Processing Systems. "
    "2014;27:3104-3112.",
    "[9] Kingma DP, Ba J. Adam: A method for stochastic optimization. "
    "International Conference on Learning Representations. 2015.",
    "[10] Radford A, Wu J, Child R, et al. Language models are unsupervised "
    "multitask learners. OpenAI Technical Report. 2019.",
    "[11] Raffel C, Shazeer N, Roberts A, et al. Exploring the limits of "
    "transfer learning with a unified text-to-text transformer. Journal of "
    "Machine Learning Research. 2020;21(140):1-67.",
    "[12] Hochreiter S, Schmidhuber J. Long short-term memory. Neural "
    "Computation. 1997;9(8):1735-1780. doi:10.1162/neco.1997.9.8.1735.",
    "[13] Bahdanau D, Cho K, Bengio Y. Neural machine translation by jointly "
    "learning to align and translate. International Conference on Learning "
    "Representations. 2015.",
    "[14] Pennington J, Socher R, Manning CD. GloVe: Global vectors for word "
    "representation. Proceedings of EMNLP. 2014:1532-1543. "
    "doi:10.3115/v1/D14-1162.",
    "[15] Rogers A, Kovaleva O, Rumshisky A. A primer in BERTology: What we "
    "know about how BERT works. Transactions of the Association for "
    "Computational Linguistics. 2020;8:842-866. doi:10.1162/tacl_a_00349.",
    "[16] Karpukhin V, Oguz B, Min S, et al. Dense passage retrieval for "
    "open-domain question answering. Proceedings of EMNLP. 2020:6769-6781. "
    "doi:10.18653/v1/2020.emnlp-main.550.",
]

AY_REFS = [
    "Brown, T. B., Mann, B., Ryder, N., Subbiah, M., Kaplan, J., et al. "
    "(2020). Language models are few-shot learners. Advances in Neural "
    "Information Processing Systems, 33, 1877-1901.",
    "Devlin, J., Chang, M.-W., Lee, K., & Toutanova, K. (2019). BERT: "
    "Pre-training of deep bidirectional transformers for language "
    "understanding. Proceedings of NAACL-HLT, 4171-4186.",
    "He, K., Zhang, X., Ren, S., & Sun, J. (2016). Deep residual learning "
    "for image recognition. Proceedings of CVPR, 770-778.",
    "Hochreiter, S., & Schmidhuber, J. (1997). Long short-term memory. "
    "Neural Computation, 9(8), 1735-1780.",
    "Ji, Z., Lee, N., Frieske, R., Yu, T., Su, D., et al. (2023). Survey of "
    "hallucination in natural language generation. ACM Computing Surveys, "
    "55(12), 1-38.",
    "Kingma, D. P., & Ba, J. (2015). Adam: A method for stochastic "
    "optimization. Proceedings of ICLR.",
    "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., et al. "
    "(2020). Retrieval-augmented generation for knowledge-intensive NLP "
    "tasks. Advances in Neural Information Processing Systems, 33, 9459-9474.",
    "Mikolov, T., Chen, K., Corrado, G., & Dean, J. (2013a). Efficient "
    "estimation of word representations in vector space. arXiv:1301.3781.",
    "Mikolov, T., Sutskever, I., Chen, K., Corrado, G., & Dean, J. (2013b). "
    "Distributed representations of words and phrases and their "
    "compositionality. Advances in Neural Information Processing Systems, 26.",
    "Vaswani, A., Shazeer, N., Parmar, N., Uszkoreit, J., Jones, L., et al. "
    "(2017). Attention is all you need. Advances in Neural Information "
    "Processing Systems, 30, 5998-6008.",
]

ZH_REFS = [
    "[1] 刘知远, 孙茂松, 林衍凯, 等. 知识表示学习研究进展[J]. 计算机研究与发展, "
    "2016, 53(2): 247-261.",
    "[2] 宗成庆. 统计自然语言处理[M]. 2版. 北京: 清华大学出版社, 2013.",
    "[3] 车万翔, 窦志成, 冯岩松, 等. 大模型时代的自然语言处理: 挑战、机遇与"
    "发展[J]. 中国科学: 信息科学, 2023, 53(9): 1645-1687.",
    "[4] 张钹, 朱军, 苏航. 迈向第三代人工智能[J]. 中国科学: 信息科学, 2020, "
    "50(9): 1281-1302.",
    "[5] 邱锡鹏. 神经网络与深度学习[M]. 北京: 机械工业出版社, 2020.",
    "[6] 周志华. 机器学习[M]. 北京: 清华大学出版社, 2016.",
    "[7] 冯志伟. 机器翻译研究[M]. 北京: 中国对外翻译出版公司, 2004.",
    "[8] 李德毅, 杜鹃. 不确定性人工智能[M]. 2版. 北京: 国防工业出版社, 2014.",
]

ZOTERO_REFS = [
    "[1] Vaswani A, Shazeer N, Parmar N, et al. Attention is all you need. "
    "Advances in Neural Information Processing Systems. 2017;30:5998-6008.",
    "[2] Devlin J, Chang M-W, Lee K, Toutanova K. BERT: Pre-training of deep "
    "bidirectional transformers for language understanding. Proceedings of "
    "NAACL-HLT. 2019:4171-4186.",
    "[3] Brown TB, Mann B, Ryder N, et al. Language models are few-shot "
    "learners. Advances in Neural Information Processing Systems. "
    "2020;33:1877-1901.",
    "[4] Ji Z, Lee N, Frieske R, et al. Survey of hallucination in natural "
    "language generation. ACM Computing Surveys. 2023;55(12):1-38.",
    "[5] Lewis P, Perez E, Piktus A, et al. Retrieval-augmented generation "
    "for knowledge-intensive NLP tasks. Advances in Neural Information "
    "Processing Systems. 2020;33:9459-9474.",
    "[6] Kingma DP, Ba J. Adam: A method for stochastic optimization. "
    "International Conference on Learning Representations. 2015.",
]

ZOTERO_ITEMS = {
    "r1": zotero_item("Attention is all you need", [("Vaswani", "Ashish"),
                      ("Shazeer", "Noam")], 2017, doi="10.48550/arXiv.1706.03762"),
    "r2": zotero_item("BERT: Pre-training of deep bidirectional transformers "
                      "for language understanding",
                      [("Devlin", "Jacob"), ("Chang", "Ming-Wei")], 2019,
                      doi="10.18653/v1/N19-1423"),
    "r3": zotero_item("Language models are few-shot learners",
                      [("Brown", "Tom"), ("Mann", "Benjamin")], 2020),
    "r4": zotero_item("Survey of hallucination in natural language generation",
                      [("Ji", "Ziwei"), ("Lee", "Nayeon")], 2023,
                      doi="10.1145/3571730"),
    "r5": zotero_item("Retrieval-augmented generation for knowledge-intensive "
                      "NLP tasks", [("Lewis", "Patrick"), ("Perez", "Ethan")], 2020),
    "r6": zotero_item("Adam: A method for stochastic optimization",
                      [("Kingma", "Diederik"), ("Ba", "Jimmy")], 2015,
                      item_type="paper-conference"),
}


def save(doc, name: str, expected: dict):
    path = HERE / name
    doc.save(path)
    expected["filename"] = name
    (HERE / (name.replace(".docx", ".expected.json"))).write_text(
        json.dumps(expected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    # sanity: it is a zip with a document.xml
    with zipfile.ZipFile(path) as zf:
        assert "word/document.xml" in zf.namelist()
    print(f"wrote {name} ({len(expected['markers'])} markers, "
          f"{expected['reference_count']} refs)")


# ---------------------------------------------------------------- fixtures


def make_numeric_en():
    doc = Document()
    doc.add_paragraph("Citation Behaviour in Neural Models", style="Title")
    doc.add_paragraph("A short study manuscript for parser testing.")

    add_heading(doc, "1 Introduction")
    add_para(doc, [
        ("Transformer architectures ", {}),
        ("[1", {"b": True}), (",3", {"i": True}), ("]", {}),
        (" have reshaped natural language processing, and pre-trained models ", {}),
        ("[2]", {}),
        (" achieve strong benchmark results.", {}),
    ])
    add_para(doc, [
        ("Retrieval-augmented generation ", {}),
        ("[4]", {}),
        (" and dense retrieval ", {}),
        ("[16]", {}),
        (" ground responses in external evidence, while hallucination remains a "
         "central concern ", {}),
        ("[5]", {}),
        (".", {}),
    ])

    add_heading(doc, "2 Related Work")
    add_para(doc, [
        ("Early sequence models relied on recurrent networks ", {}),
        ("[8]", {}),
        (" and long short-term memory ", {}),
        ("[12]", {}),
        ("; attention mechanisms ", {}),
        ("[1,13]", {}),
        (" improved alignment.", {}),
    ])
    add_para(doc, [
        ("Word representations evolved from word2vec ", {}),
        ("[7]", {}),
        (" to contextual embeddings ", {}),
        ("[2]", {}),
        (", and optimization is typically handled by Adam ", {}),
        ("[9]", {}),
        (".", {}),
    ])
    add_para(doc, [
        ("Surveys of BERTology ", {}),
        ("[15]", {}),
        (" and transfer learning ", {}),
        ("[11]", {}),
        (" summarize fine-tuning practice ", {}),
        ("[2,4-6,9]", {}),
        (".", {}),
    ])

    add_heading(doc, "3 Method")
    add_para(doc, [
        ("We build on the Transformer encoder ", {}),
        ("[1]", {}),
        (" and initialize from BERT ", {}),
        ("[2]", {}),
        (". Following earlier work ", {}),
        ("[2-5]", {}),
        (", we fine-tune with residual connections ", {}),
        ("[6]", {}),
        (" and dropout.", {}),
    ])
    add_para(doc, [
        ("The full pipeline ", {}),
        ("[1,4,16]", {"b": True}),
        (" combines\tretrieval and generation in one stage.", {}),
    ])

    add_heading(doc, "4 Experiments")
    add_para(doc, [
        ("We evaluate on three benchmarks ", {}),
        ("[2,4-6,9]", {}),
        (" using standard metrics ", {}),
        ("[10-11]", {}),
        (".", {}),
    ])
    table: Table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("Setting")
    table.cell(0, 1).paragraphs[0].add_run("Result")
    table.cell(1, 0).paragraphs[0].add_run("baseline GloVe ")
    r = table.cell(1, 0).paragraphs[0].add_run("[14]")
    table.cell(1, 1).paragraphs[0].add_run("reported in ")
    r2 = table.cell(1, 1).paragraphs[0].add_run("8")
    r2.font.superscript = True

    add_heading(doc, "5 Discussion")
    add_para(doc, [
        ("Our findings align with prior reports ", {}),
        ("[15]", {}),
        (" although coverage differs ", {}),
        ("[3,10]", {}),
        (".", {}),
    ])

    add_heading(doc, "6 Conclusion")
    add_para(doc, [
        ("We presented a compact pipeline ", {}),
        ("[1]", {}),
        (" that performs on par with larger systems ", {}),
        ("[3,10,11]", {}),
        (".", {}),
    ])

    add_heading(doc, "References")
    for raw in EN_REFS:
        add_para(doc, [(raw, {})])

    markers = [
        {"raw": "[1,3]", "ref_ids": ["r1", "r3"]},
        {"raw": "[2]", "ref_ids": ["r2"]},
        {"raw": "[4]", "ref_ids": ["r4"]},
        {"raw": "[16]", "ref_ids": ["r16"]},
        {"raw": "[5]", "ref_ids": ["r5"]},
        {"raw": "[8]", "ref_ids": ["r8"]},
        {"raw": "[12]", "ref_ids": ["r12"]},
        {"raw": "[1,13]", "ref_ids": ["r1", "r13"]},
        {"raw": "[7]", "ref_ids": ["r7"]},
        {"raw": "[2]", "ref_ids": ["r2"]},
        {"raw": "[9]", "ref_ids": ["r9"]},
        {"raw": "[15]", "ref_ids": ["r15"]},
        {"raw": "[11]", "ref_ids": ["r11"]},
        {"raw": "[2,4-6,9]", "ref_ids": ["r2", "r4", "r5", "r6", "r9"]},
        {"raw": "[1]", "ref_ids": ["r1"]},
        {"raw": "[2]", "ref_ids": ["r2"]},
        {"raw": "[2-5]", "ref_ids": ["r2", "r3", "r4", "r5"]},
        {"raw": "[6]", "ref_ids": ["r6"]},
        {"raw": "[1,4,16]", "ref_ids": ["r1", "r4", "r16"]},
        {"raw": "[2,4-6,9]", "ref_ids": ["r2", "r4", "r5", "r6", "r9"]},
        {"raw": "[10-11]", "ref_ids": ["r10", "r11"]},
        {"raw": "[14]", "ref_ids": ["r14"]},  # inside the table cell
        {"raw": "8", "ref_ids": ["r8"]},      # superscript, table cell
        {"raw": "[15]", "ref_ids": ["r15"]},
        {"raw": "[3,10]", "ref_ids": ["r3", "r10"]},
        {"raw": "[1]", "ref_ids": ["r1"]},
        {"raw": "[3,10,11]", "ref_ids": ["r3", "r10", "r11"]},
    ]
    save(doc, "numeric_en.docx", {
        "citation_style": "numeric",
        "managed_by": None,
        "reference_count": 16,
        "sections": [
            {"title": "1 Introduction", "canonical": "intro"},
            {"title": "2 Related Work", "canonical": "related"},
            {"title": "3 Method", "canonical": "method"},
            {"title": "4 Experiments", "canonical": "experiment"},
            {"title": "5 Discussion", "canonical": "discussion"},
            {"title": "6 Conclusion", "canonical": "conclusion"},
            {"title": "References", "canonical": "other"},
        ],
        "markers": markers,
    })


def make_authoryear_en():
    doc = Document()
    doc.add_paragraph("Retrieval and Reasoning", style="Title")

    add_heading(doc, "1 Introduction")
    add_para(doc, [
        ("Attention-based architectures ", {}),
        ("(Vaswani ", {"i": True}), ("et al.", {}), (", 2017)", {}),
        (" underpin modern language models ", {}),
        ("(Devlin et al., 2019; Brown et al., 2020)", {}),
        (".", {}),
    ])
    add_para(doc, [
        ("Optimization relies on Adam ", {}),
        ("(Kingma & Ba, 2015)", {}),
        (", a descendant of recurrent training practice ", {}),
        ("(Hochreiter & Schmidhuber, 1997)", {}),
        (".", {}),
    ])

    add_heading(doc, "2 Related Work")
    add_para(doc, [
        ("Mikolov et al. (2013a) introduced efficient embeddings, and ", {}),
        ("Mikolov et al. (2013b) extended them to phrases; retrieval-based "
         "approaches ", {}),
        ("(Lewis et al., 2020)", {}),
        (" connect models to corpora.", {}),
    ])
    add_para(doc, [
        ("Residual learning ", {}),
        ("(He et al., 2016)", {}),
        (" stabilizes depth, while hallucination surveys ", {}),
        ("(Ji et al., 2023)", {}),
        (" catalogue failure modes.", {}),
    ])

    add_heading(doc, "3 Method")
    add_para(doc, [
        ("Following ", {}),
        ("Devlin et al. (2019)", {}),
        (", we fine-tune a masked language model; Vaswani et al. (2017) provide "
         "the encoder stack.", {}),
    ])

    add_heading(doc, "References")
    for raw in AY_REFS:
        add_para(doc, [(raw, {})])

    markers = [
        {"raw": "(Vaswani et al., 2017)", "ref_ids": ["r10"]},
        {"raw": "(Devlin et al., 2019; Brown et al., 2020)",
         "ref_ids": ["r2", "r1"]},
        {"raw": "(Kingma & Ba, 2015)", "ref_ids": ["r6"]},
        {"raw": "(Hochreiter & Schmidhuber, 1997)", "ref_ids": ["r4"]},
        {"raw": "Mikolov et al. (2013a)", "ref_ids": ["r8"]},
        {"raw": "Mikolov et al. (2013b)", "ref_ids": ["r9"]},
        {"raw": "(Lewis et al., 2020)", "ref_ids": ["r7"]},
        {"raw": "(He et al., 2016)", "ref_ids": ["r3"]},
        {"raw": "(Ji et al., 2023)", "ref_ids": ["r5"]},
        {"raw": "Devlin et al. (2019)", "ref_ids": ["r2"]},
        {"raw": "Vaswani et al. (2017)", "ref_ids": ["r10"]},
    ]
    save(doc, "authoryear_en.docx", {
        "citation_style": "author_year",
        "managed_by": None,
        "reference_count": 10,
        "sections": [
            {"title": "1 Introduction", "canonical": "intro"},
            {"title": "2 Related Work", "canonical": "related"},
            {"title": "3 Method", "canonical": "method"},
            {"title": "References", "canonical": "other"},
        ],
        "markers": markers,
    })


def make_gbt_zh():
    doc = Document()
    inject_localized_styles(doc)

    doc.add_paragraph("面向学术写作的引用一致性检测", style="Title")

    h = doc.add_paragraph("一、引言")
    set_pstyle(h, "1")
    add_para(doc, [
        ("知识表示学习", {}),
        ("［1］", {}),
        ("是自然语言处理的基础问题", {}),
        ("[2]", {}),
        ("。大模型时代的到来", {}),
        ("[3]", {}),
        ("对引用规范提出了新的要求。", {}),
    ])

    h = doc.add_paragraph("二、相关工作")
    set_pstyle(h, "1")
    add_para(doc, [
        ("第三代人工智能", {}),
        ("[4]", {}),
        ("强调知识驱动与数据驱动结合。神经网络方法", {}),
        ("[5,6]", {}),
        ("在多个任务上取得进展。", {}),
    ])

    h = doc.add_paragraph("三、研究方法")
    set_pstyle(h, "1")
    add_para(doc, [
        ("本文在文献分析基础上", {}),
        ("[1,3]", {}),
        ("构建引用检测流水线。早期机器翻译研究", {}),
        ("[7]", {}),
        ("提供了规则方法的经验。", {}),
    ])
    h2 = doc.add_paragraph("3.1 数据来源")
    set_pstyle(h2, "2")  # localized Heading 2 -> stays inside section 三
    add_para(doc, [
        ("语料来源于公开期刊", {}),
        ("[3,4,8]", {"b": True}),
        ("，覆盖2016年至2023年的代表性工作。", {}),
    ])

    h = doc.add_paragraph("四、结论")
    set_pstyle(h, "1")
    add_para(doc, [
        ("实验表明该方法可以有效识别引用异常", {}),
        ("[2-4]", {}),
        ("。", {}),
    ])

    h = doc.add_paragraph("参考文献")
    set_pstyle(h, "1")
    for raw in ZH_REFS:
        add_para(doc, [(raw, {})])

    markers = [
        {"raw": "［1］", "ref_ids": ["r1"]},
        {"raw": "[2]", "ref_ids": ["r2"]},
        {"raw": "[3]", "ref_ids": ["r3"]},
        {"raw": "[4]", "ref_ids": ["r4"]},
        {"raw": "[5,6]", "ref_ids": ["r5", "r6"]},
        {"raw": "[1,3]", "ref_ids": ["r1", "r3"]},
        {"raw": "[7]", "ref_ids": ["r7"]},
        {"raw": "[3,4,8]", "ref_ids": ["r3", "r4", "r8"]},
        {"raw": "[2-4]", "ref_ids": ["r2", "r3", "r4"]},
    ]
    save(doc, "gbt_zh.docx", {
        "citation_style": "numeric",
        "managed_by": None,
        "reference_count": 8,
        "sections": [
            {"title": "一、引言", "canonical": "intro"},
            {"title": "二、相关工作", "canonical": "related"},
            {"title": "三、研究方法", "canonical": "method"},
            {"title": "四、结论", "canonical": "conclusion"},
            {"title": "参考文献", "canonical": "other"},
        ],
        "markers": markers,
    })


def make_zotero_numeric():
    doc = Document()
    doc.add_paragraph("Field-Managed Citation Study", style="Title")

    add_heading(doc, "1 Introduction")
    p = add_para(doc, [("Transformers ", {}), ])
    add_complex_field(p, zotero_instr([ZOTERO_ITEMS["r1"]], "[1]"), "[1]")
    p._p.append(etree.fromstring(f"<w:r {ns}><w:t> and BERT </w:t></w:r>"))
    add_complex_field(p, zotero_instr([ZOTERO_ITEMS["r2"]], "[2]"), "[2]")
    p._p.append(etree.fromstring(
        f"<w:r {ns}><w:t> dominate recent benchmarks.</w:t></w:r>"))

    add_heading(doc, "2 Related Work")
    p = add_para(doc, [("Several surveys ", {})])
    add_complex_field(
        p,
        zotero_instr([ZOTERO_ITEMS["r3"], ZOTERO_ITEMS["r4"]], "[3], [4]"),
        "[3], [4]",
    )
    p._p.append(etree.fromstring(
        f"<w:r {ns}><w:t> discuss open problems, while retrieval methods </w:t></w:r>"))
    add_complex_field(p, zotero_instr([ZOTERO_ITEMS["r5"]], "[5]"), "[5]")
    p._p.append(etree.fromstring(
        f"<w:r {ns}><w:t> ground generation; optimizers such as Adam </w:t></w:r>"))
    add_complex_field(p, zotero_instr([ZOTERO_ITEMS["r6"]], "[6]"), "[6]")
    p._p.append(etree.fromstring(
        f"<w:r {ns}><w:t> remain standard.</w:t></w:r>"))

    add_heading(doc, "3 Method")
    p = add_para(doc, [("We follow the setup of ", {})])
    add_complex_field(p, zotero_instr([ZOTERO_ITEMS["r2"]], "[2]"), "[2]")
    p._p.append(etree.fromstring(f"<w:r {ns}><w:t> closely.</w:t></w:r>"))

    add_heading(doc, "References")
    # ZOTERO_BIBL field: begin/instr/separate in the first entry paragraph,
    # end fldChar in the last entry paragraph.
    bibl_instr = ' ZOTERO_BIBL {"custom":[]} CSL_BIBLIOGRAPHY '
    first = add_para(doc, [])
    for frag in (
        '<w:fldChar w:fldCharType="begin"/>',
        f'<w:instrText xml:space="preserve">{_esc(bibl_instr)}</w:instrText>',
        '<w:fldChar w:fldCharType="separate"/>',
        f"<w:t>{_esc(ZOTERO_REFS[0])}</w:t>",
    ):
        first._p.append(etree.fromstring(f"<w:r {ns}>{frag}</w:r>"))
    for raw in ZOTERO_REFS[1:-1]:
        add_para(doc, [(raw, {})])
    last = add_para(doc, [(ZOTERO_REFS[-1], {})])
    last._p.append(etree.fromstring(
        f'<w:r {ns}><w:fldChar w:fldCharType="end"/></w:r>'))

    markers = [
        {"raw": "[1]", "ref_ids": ["r1"]},
        {"raw": "[2]", "ref_ids": ["r2"]},
        {"raw": "[3], [4]", "ref_ids": ["r3", "r4"]},
        {"raw": "[5]", "ref_ids": ["r5"]},
        {"raw": "[6]", "ref_ids": ["r6"]},
        {"raw": "[2]", "ref_ids": ["r2"]},
    ]
    save(doc, "zotero_numeric.docx", {
        "citation_style": "numeric",
        "managed_by": "zotero",
        "reference_count": 6,
        "sections": [
            {"title": "1 Introduction", "canonical": "intro"},
            {"title": "2 Related Work", "canonical": "related"},
            {"title": "3 Method", "canonical": "method"},
            {"title": "References", "canonical": "other"},
        ],
        "markers": markers,
    })


if __name__ == "__main__":
    make_numeric_en()
    make_authoryear_en()
    make_gbt_zh()
    make_zotero_numeric()
