"""Word tables, embedded images and captions for the reading view."""

import io

from docx import Document as Docx

from citecheck.parse import parse_document
from citecheck.parse.assets import asset_path

_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000b49444154789c6360000200000500017a5eab3f0000"
    "000049454e44ae426082")


def test_tables_images_captions(tmp_path):
    doc = Docx()
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Table 1 shows the scores [1].")
    doc.add_paragraph("Table 1: Main results.", style="Caption")
    t = doc.add_table(rows=3, cols=3)
    t.cell(0, 1).merge(t.cell(0, 2))
    t.cell(0, 1).text = "BLEU"
    for c, v in enumerate(["Model", "EN-DE", "EN-FR"]):
        t.cell(1, c).text = v
    for c, v in enumerate(["Base", "27.3", "38.1"]):
        t.cell(2, c).text = v
    doc.add_picture(io.BytesIO(_PNG))
    doc.add_paragraph("Figure 2. The architecture.")
    path = tmp_path / "paper.docx"
    doc.save(path)

    d = parse_document(path)
    cells = [(p.cell.row, p.cell.col, p.cell.colspan, p.cell.header, p.text)
             for p in d.paragraphs if p.cell and p.text]
    assert cells == [
        (0, 1, 2, True, "BLEU"),
        (1, 0, 1, False, "Model"), (1, 1, 1, False, "EN-DE"),
        (1, 2, 1, False, "EN-FR"),
        (2, 0, 1, False, "Base"), (2, 1, 1, False, "27.3"),
        (2, 2, 1, False, "38.1"),
    ]
    caps = [p.text for p in d.paragraphs if p.caption]
    assert caps == ["Table 1: Main results.", "Figure 2. The architecture."]
    img = next(p for p in d.paragraphs if p.images)
    assert img.text == "" and img.cell is None
    got = asset_path(d.document.assets, img.images[0])
    assert got is not None and got.read_bytes() == _PNG
