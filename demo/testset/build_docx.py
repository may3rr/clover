"""F33: rebuild the five Chinese eval manuscripts as readable Word documents.

The earlier version (F8) pulled text lines out of the journal PDFs and kept
only the reference list and [n] markers intact; headings merged into body
text, tables flattened into prose ("表 1" turned into the marker "表[1]"),
figures vanished. This version rebuilds the document structure:

- PDF sources (CNKI downloads, the AAS publisher PDF): characters are read
  with their own size and baseline (``rawdict``), grouped into visual rows,
  laid out by column (full-width items such as wide figures split the page
  into bands; inside a band the left column is read before the right one).
  Superscript citation numbers are attached to the row below them and
  written as superscript ``[n]`` runs. Numbered headings are accepted only
  when their number continues the previous heading's numbering. Figures are
  cropped from the page above their caption; tables are rebuilt as Word
  tables from the positions of their cell text.
- The 软件学报 paper comes from the journal's HTML full text, which already
  carries headings, figures, tables and superscript citations.

Reference lists keep their [n] labels. Raw sources (PDF, HTML, images) are
not committed; see README.md.

Usage:
    python demo/testset/build_docx.py            # all five
    python demo/testset/build_docx.py tod kg     # by key
Outputs: zh_papers/<dir>/manuscript.docx
"""

from __future__ import annotations

import io
import json
import re
import statistics
import sys
import unicodedata
import urllib.request
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import fitz
import lxml.html
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent


# Chinese punctuation and circled numbers stay as typeset: a real
# manuscript has "，．；（）" and "①", not their ASCII folds. Full-width
# letters, digits and brackets (CNKI sets Latin text full-width) fold.
_KEEP = set("，．；：！？（）、。“”‘’《》【】—…·")


def nk(s: str) -> str:
    return "".join(
        c if c in _KEEP or 0x2460 <= ord(c) <= 0x24FF or 0x2776 <= ord(c) <= 0x2793
        else unicodedata.normalize("NFKC", c) for c in s)


_TIMES = fitz.Font("tiro")


def _advance_end(c: "Ch") -> float:
    """Right edge of a glyph. CNKI fonts carry no usable widths (every bbox
    is a full em), so latin glyphs use Times advance widths from their
    origin instead."""
    if c.c.isascii() and c.c.isprintable():
        return c.x0 + min(c.x1 - c.x0, _TIMES.glyph_advance(ord(c.c)) * c.size * 1.05)
    return c.x1


def _latin(c: str) -> bool:
    """ASCII glyph that takes part in word spacing (not CJK, not
    full-width punctuation folded to ASCII by NFKC next to CJK)."""
    return c.isascii() and c.isprintable() and c != " "


def is_cjk(c: str) -> bool:
    return "㐀" <= c <= "鿿" or "豈" <= c <= "﫿"


def is_cjk_punct(c: str) -> bool:
    return c in "，．。、；：？！“”‘’（）《》【】—…·" or "　" <= c <= "〿"


# ---------------------------------------------------------------------------
# document blocks, shared by both engines
# ---------------------------------------------------------------------------

@dataclass
class Block:
    kind: str  # title | front | para | heading | caption | figure | table | refhead | ref | formula
    runs: list[tuple[str, bool]] = field(default_factory=list)  # (text, superscript)
    level: int = 1
    image: bytes | None = None
    width_pt: float = 0.0
    rows: list[list[str]] | None = None
    center: bool = False

    @property
    def text(self) -> str:
        return "".join(t for t, _ in self.runs)


def plain(kind: str, text: str, **kw) -> Block:
    return Block(kind, [(text, False)], **kw)


# ---------------------------------------------------------------------------
# PDF engine
# ---------------------------------------------------------------------------

@dataclass
class Ch:
    c: str
    x0: float
    x1: float
    y0: float
    y1: float
    base: float
    size: float
    sup: bool = False
    minor: bool = False  # inline sub/superscript of a formula
    space: bool = False  # the PDF has a space character before this glyph


@dataclass
class Row:
    chars: list[Ch]
    page: int = 0
    kind: str = "text"  # text | caption
    centred: bool = False  # set during paragraph assembly

    @property
    def x0(self): return min(c.x0 for c in self.chars)
    @property
    def x1(self): return max(c.x1 for c in self.chars)
    @property
    def y0(self):
        main = [c.y0 for c in self.chars if not (c.sup or c.minor)]
        return min(main) if main else min(c.y0 for c in self.chars)
    @property
    def y1(self): return max(c.y1 for c in self.chars)
    @property
    def size(self):
        s = [c.size for c in self.chars if not (c.sup or c.minor)]
        return statistics.median(s) if s else statistics.median(c.size for c in self.chars)
    @property
    def base(self):
        b = [c.base for c in self.chars if not (c.sup or c.minor)]
        return statistics.median(b) if b else self.chars[0].base

    def runs(self, markers: bool = True) -> list[tuple[str, bool]]:
        out: list[tuple[str, bool]] = []
        chars = sorted(self.chars, key=lambda c: c.x0)
        # word gaps between latin glyphs: residual after the glyph's
        # advance, measured against this row's typical in-word residual
        # (the embedded font is never exactly Times)
        resid = [b.x0 - _advance_end(a) for a, b in zip(chars, chars[1:])
                 if _latin(a.c) and _latin(b.c) and not a.sup and not b.sup]
        base = statistics.median(resid) if len(resid) >= 6 else 0.0
        # rows whose PDF text carries its own spaces (reference lists) need
        # no guessing; CNKI abstracts and titles carry none
        spaced = any(c.space and _latin(c.c) for c in chars)
        prev: Ch | None = None
        for c in chars:
            sup = c.sup and markers
            sep = ""
            if prev is not None and not c.sup and not prev.sup and is_cjk(prev.c) \
                    and is_cjk(c.c) and c.x0 - prev.x1 > 0.6 * c.size:
                sep = " "  # names set apart in author lines
            elif spaced and prev is not None and not c.sup and not prev.sup:
                if c.space and not (is_cjk(prev.c) and is_cjk(c.c)):
                    sep = " "
            elif prev is not None and not c.sup and not prev.sup \
                    and _latin(prev.c) and _latin(c.c):
                gap = c.x0 - _advance_end(prev) - max(base, 0.0)
                if gap > 0.2 * min(c.size, prev.size) and c.c not in ",.;:?!)]-" \
                        and prev.c not in "([-/":
                    sep = " "
            if out and out[-1][1] == sup:
                out[-1] = (out[-1][0] + sep + c.c, sup)
            else:
                if sep and out:
                    out[-1] = (out[-1][0] + sep, out[-1][1])
                out.append((c.c, sup))
            prev = c
        return [(_marker_text(t) if s else t, s) for t, s in out]

    @property
    def text(self) -> str:
        return "".join(t for t, _ in self.runs())

    @property
    def squashed(self) -> str:
        return re.sub(r"\s+", "", self.text)


_PUNCT = set(",.;:!?、。，．；：！？·()（）“”‘’《》\"'")
_MARKER_CHARS = re.compile(r"^[\[\]\d,，、\-–—~～−\s]+$")


def _marker_text(t: str) -> str:
    t = t.replace("，", ",").replace("−", "-").replace("–", "-").replace(" ", "")
    if not t.startswith("["):
        t = "[" + t
    if not t.endswith("]"):
        t = t + "]"
    return t


def page_chars(page: fitz.Page) -> list[Ch]:
    out = []
    for b in page.get_text("rawdict")["blocks"]:
        if b.get("type") != 0:
            continue
        for ln in b["lines"]:
            if abs(ln["dir"][1]) > 0.1:
                continue
            gap = False  # a space character came before the next glyph
            for s in ln["spans"]:
                for ch in s["chars"]:
                    c = nk(ch["c"])
                    if not c.strip():
                        gap = True
                        continue
                    x0, y0, x1, y1 = ch["bbox"]
                    # NFKC can expand one glyph (ﬁ -> fi): spread it evenly
                    w = (x1 - x0) / len(c)
                    for k, cc in enumerate(c):
                        out.append(Ch(cc, x0 + k * w, x0 + (k + 1) * w, y0, y1,
                                      ch["origin"][1], s["size"], space=gap and k == 0))
                    gap = False
    return out


def find_gutter(doc: fitz.Document) -> float:
    """x of the gap between the two text columns (fewest body characters)."""
    W = doc[0].rect.width
    lo, hi = int(W * 0.4), int(W * 0.6)
    cover = Counter()
    for pno in range(1, min(len(doc), 8)):
        for c in page_chars(doc[pno]):
            for x in range(max(lo, int(c.x0) + 1), min(hi, int(c.x1))):
                cover[x] += 1
    xs = list(range(lo, hi))
    best = min(cover[x] for x in xs)
    run = [x for x in xs if cover[x] == best]
    # centre of the longest run of minimal coverage
    groups, cur = [], [run[0]]
    for x in run[1:]:
        if x == cur[-1] + 1:
            cur.append(x)
        else:
            groups.append(cur)
            cur = [x]
    groups.append(cur)
    g = max(groups, key=len)
    return (g[0] + g[-1]) / 2


def build_rows(chars: list[Ch], gutter: float, body: float) -> list[Row]:
    def side(x0: float, x1: float) -> str:
        return "L" if x1 <= gutter + 1 else "R" if x0 >= gutter - 1 else "F"

    # 1. baseline lines; a glyph joins an open line on its own side of the
    #    gutter first (a heading number a point off its title's baseline
    #    must not join the other column's line)
    punct = [c for c in chars if c.c in _PUNCT]
    chars = sorted((c for c in chars if c.c not in _PUNCT), key=lambda c: (c.base, c.x0))
    lines: list[dict] = []
    for c in chars:
        cs = side(c.x0, c.x1)
        cands = [ln for ln in lines[-16:]
                 if abs(c.base - ln["base"]) <= 0.25 * min(c.size, ln["size"]) + 0.6]
        same = [ln for ln in cands if ln["side"] in (cs, "F")]
        near = [ln for ln in cands
                if min(abs(c.x0 - ln["x1"]), abs(ln["x0"] - c.x1)) < 1.5 * c.size]
        pick = min(same or near, key=lambda ln: abs(c.base - ln["base"]), default=None)
        if pick is None:
            lines.append(dict(base=c.base, size=c.size, side=cs, x0=c.x0, x1=c.x1, chars=[c]))
        else:
            pick["chars"].append(c)
            pick["x0"], pick["x1"] = min(pick["x0"], c.x0), max(pick["x1"], c.x1)
            if pick["side"] != cs:
                pick["side"] = "F"

    # punctuation glyphs sit in their own font with an offset baseline:
    # they join the closest line whose span covers them
    for c in punct:
        cands = [ln for ln in lines if abs(c.base - ln["base"]) <= 0.7 * ln["size"]
                 and c.size >= 0.8 * ln["size"]
                 and ln["x0"] - 1.2 * ln["size"] <= c.x0 <= ln["x1"] + 1.2 * ln["size"]]
        if cands:
            ln = min(cands, key=lambda ln: abs(c.base - ln["base"]))
            ln["chars"].append(c)
            ln["x0"], ln["x1"] = min(ln["x0"], c.x0), max(ln["x1"], c.x1)
        else:
            lines.append(dict(base=c.base, size=c.size, side=side(c.x0, c.x1),
                              x0=c.x0, x1=c.x1, chars=[c]))

    # 2. split lines into rows at the gutter and at wide gaps (table cells)
    rows: list[Row] = []
    for ln in lines:
        cs = sorted(ln["chars"], key=lambda c: c.x0)
        spans_gutter = any(c.x0 < gutter - 1 and c.x1 > gutter + 1 for c in cs)
        seg = [cs[0]]
        for prev, c in zip(cs, cs[1:]):
            gap = c.x0 - prev.x1
            size = max(c.size, prev.size)
            split = gap > max(2.6 * size, 16)
            if not spans_gutter and gap > 0.55 * size \
                    and prev.x1 <= gutter + 2 and c.x0 >= gutter - 2:
                split = True
            if split:
                rows.append(Row(seg))
                seg = [c]
            else:
                seg.append(c)
        rows.append(Row(seg))

    # 3. superscripts / subscripts: short small rows just above / below a
    #    larger row join it; citation numbers above it become marker runs
    def small(r: Row) -> bool:
        return len(r.chars) <= 16 and r.size <= 0.8 * body

    big = [r for r in rows if not small(r)]
    out = list(big)
    for s in (r for r in rows if small(r)):
        host = None
        for r in big:
            if r.size < s.size * 1.2:
                continue
            d = r.base - s.base
            if not (-0.5 * r.size <= d <= 0.75 * r.size) or abs(d) < 0.4:
                continue
            if s.x1 < r.x0 - 4 or s.x0 > r.x1 + 4:
                continue
            host = r
            break
        if host is None:
            out.append(s)
            continue
        digits = "".join(c.c for c in s.chars)
        marker = d >= 0.4 and _MARKER_CHARS.match(digits) and re.search(r"[1-9]", digits)
        for c in s.chars:
            c.sup = bool(marker)
            c.minor = not marker
        host.chars.extend(s.chars)
    return out


def graphics_clusters(page: fitz.Page):
    """(figure clusters, horizontal rules) as fitz.Rect lists."""
    rects = []
    for info in page.get_image_info():
        r = fitz.Rect(info["bbox"]) & page.rect
        if r.width * r.height > 300:
            rects.append(r)
    rules = []
    for d in page.get_drawings():
        r = fitz.Rect(d["rect"])
        if r.height < 2.2 and r.width > 30:
            rules.append(r)
            continue
        if r.width < 2.2 and r.height > 8:
            rules.append(r)  # vertical rule
            continue
        if r.width < 0.5 and r.height < 0.5:
            continue
        rects.append(r)
    # merge overlapping / touching rects
    merged = True
    rects = [fitz.Rect(r.x0 - 3, r.y0 - 3, r.x1 + 3, r.y1 + 3) for r in rects]
    while merged:
        merged = False
        out: list[fitz.Rect] = []
        for r in rects:
            for o in out:
                if o.intersects(r):
                    o |= r
                    merged = True
                    break
            else:
                out.append(fitz.Rect(r))
        rects = out
    figs = [r for r in rects if r.width > 40 and r.height > 25]
    return figs, rules


_FIG_CAP = re.compile(r"^图\s*\d+")
_TAB_CAP = re.compile(r"^(续)?表\s*\d+|^[(（]?续表\d*[)）]?$")
_EN_CAP = re.compile(r"^(Fig\.?|Figure|Table)\s*\d+", re.I)
_HEAD = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,3})\.?\s*([^\d\s.,，。;；:：、%)）].{0,40})$")
_UNNUMBERED = re.compile(r"^(引言|结束语|结论|结语|致谢|附录|参考文献)$")
_FORCE = re.compile(r"^(摘\s*要|关键词|关\s*键\s*词|Abstract|Key\s*words|Keywords|"
                    r"中图法分类号|中图分类号|DOI|文献标识码|文章编号|引用格式)")
_TERMINAL = tuple("。．！？.!?；;:：")
_DROP_PARA = re.compile(
    r"^(收稿日期|修回日期|基金项目|本课题得到|作者简介|通信作者|通讯作者|责任编委|"
    r"Manuscript received|Recommended by|Supported by|This work was supported|"
    r"引用格式|Citation|网络出版|DOI|中图法分类号|中图分类号|文献标识码|文章编号|"
    r"\(C\)\s*1994|©)",
    re.I)
_BIO = re.compile(
    r"(教授|副教授|博士研究生|硕士研究生|讲师|研究员|博士|硕士).{0,40}(研究方向|主要研究)"
    r"|E-?mail|born\s+in\s+\d{4}|research\s+interests?|Background\b")


def _plausible(prev: list[int], cand: list[int]) -> bool:
    if not prev:
        return cand in ([0], [1])
    if cand == prev + [1]:
        return True
    for k in range(len(prev)):
        if cand[:k] == prev[:k] and len(cand) == k + 1 and cand[k] == prev[k] + 1:
            return True
    return False


@dataclass
class Item:
    kind: str  # row | figure | table
    page: int
    rect: fitz.Rect
    side: str = "F"
    row: Row | None = None
    block: Block | None = None


def _side(x0: float, x1: float, gutter: float) -> str:
    if x1 <= gutter + 3:
        return "L"
    if x0 >= gutter - 3:
        return "R"
    return "F"


def _table_block(rows: list[Row], rules: list[fitz.Rect]) -> Block | None:
    chars = [c for r in rows for c in r.chars]
    lo = min(c.x0 for c in chars)
    hi = max(_advance_end(c) for c in chars)
    top = min(c.y0 for c in chars)
    bot = max(c.y1 for c in chars)
    inside = [r for r in rules if r.y1 >= top - 30 and r.y0 <= bot + 30
              and r.x1 >= lo - 30 and r.x0 <= hi + 30]
    vert = sorted({round((r.x0 + r.x1) / 2) for r in inside if r.width < 2.2})
    seps = [x for x in vert if lo + 3 < x < hi - 3]
    if not seps:
        # three-line table: columns are the x ranges (almost) no line covers;
        # a long URL or a wrapped cell may stray across a separator
        lns: list[list[Ch]] = []
        for c in sorted(chars, key=lambda c: (c.base, c.x0)):
            if lns and abs(lns[-1][0].base - c.base) < 0.5 * c.size:
                lns[-1].append(c)
            else:
                lns.append([c])
        cover = [0] * (int(hi - lo) + 2)
        for ln in lns:
            hit = set()
            for c in ln:
                hit.update(range(int(c.x0 - lo), int(_advance_end(c) - lo) + 1))
            for x in hit:
                if 0 <= x < len(cover):
                    cover[x] += 1
        limit = max(0, int(0.12 * len(lns)))
        run = 0
        for x, v in enumerate(cover):
            if v <= limit:
                run += 1
            else:
                if run >= 4 and x - run > 0:
                    seps.append(lo + x - run / 2)
                run = 0
    ncol = len(seps) + 1
    if ncol < 2:
        return None
    # row boundaries from horizontal rules inside the table, if any
    hlines = sorted({round(r.y0) for r in inside if r.width >= 2.2 and top - 1 < r.y0 < bot + 1}) \
        if vert else []

    def band(c: Ch) -> int:
        y = (c.y0 + c.y1) / 2
        return sum(1 for h in hlines if y > h)

    # lines by baseline, then cells by column
    lines: list[list[Ch]] = []
    for c in sorted(chars, key=lambda c: (c.base, c.x0)):
        for ln in reversed(lines[-8:]):
            if abs(ln[0].base - c.base) < 0.5 * c.size:
                ln.append(c)
                break
        else:
            lines.append([c])
    lines.sort(key=lambda ln: ln[0].base)
    grid: list[list[str]] = []
    grid_band: list[int] = []
    for ln in lines:
        cols: list[list[Ch]] = [[] for _ in range(ncol)]
        for c in ln:
            cx = (c.x0 + _advance_end(c)) / 2
            cols[sum(1 for x in seps if cx > x)].append(c)
        cells = [Row(cc).text.strip() if cc else "" for cc in cols]
        b = band(ln[0]) if hlines else -1
        cont = (hlines and grid and grid_band[-1] == b) or \
            (not hlines and grid and not cells[0] and sum(1 for c in cells if c) < ncol)
        if cont:
            for k, c in enumerate(cells):
                if c:
                    grid[-1][k] = _join(grid[-1][k], c)
        else:
            grid.append(cells)
            grid_band.append(b)
    return Block("table", rows=grid)


def _join(a: str, b: str) -> str:
    if not a:
        return b
    if not b:
        return a
    if a.endswith("-") and len(a) > 1 and a[-2].isalpha() and b[0].islower():
        return a[:-1] + b
    if (a[-1].isascii() and a[-1].isalnum() or a[-1] in ",.;:)") and b[0].isascii() \
            and (b[0].isalnum() or b[0] == "("):
        return a + " " + b
    return a + b


def _join_runs(runs: list[tuple[str, bool]], more: list[tuple[str, bool]]):
    if not runs:
        return list(more)
    if not more:
        return runs
    a, sa = runs[-1]
    b, sb = more[0]
    j = _join(a, b) if not sa and not sb else None
    if j is not None:
        return runs[:-1] + [(j, False)] + more[1:]
    return runs + more


class PdfLayout:
    def __init__(self, pdf: Path, cfg: dict):
        self.doc = fitz.open(pdf)
        self.cfg = cfg
        self.gutter = cfg.get("gutter") or find_gutter(self.doc)
        sizes = Counter()
        for pno in range(1, min(len(self.doc), 6)):
            for c in page_chars(self.doc[pno]):
                if is_cjk(c.c):
                    sizes[round(c.size, 1)] += 1
        self.body = max(sizes, key=sizes.get)
        self.dropped: list[str] = []

    # -- page -> ordered items ------------------------------------------
    def page_items(self, pno: int) -> list[Item]:
        page = self.doc[pno]
        H = page.rect.height
        top = self.cfg.get("top", 70)
        bottom = H - self.cfg.get("bottom", 40)
        rows = []
        for r in build_rows(page_chars(page), self.gutter, self.body):
            if r.base < top or r.y0 > bottom:
                if r.text.strip():
                    self.dropped.append(f"p{pno} band: {r.text[:60]}")
                continue
            r.page = pno
            rows.append(r)
        figs, rules = graphics_clusters(page)
        items: list[Item] = []
        used: set[int] = set()

        caps = [r for r in rows if self._is_caption(r)]
        for cap in caps:
            sq = cap.squashed
            if _FIG_CAP.match(sq):
                region = self._figure_region(cap, figs, rows)
                if region is None:
                    self.dropped.append(f"p{pno} figure without graphics: {sq[:30]}")
                    continue
                for i, r in enumerate(rows):
                    if r is cap:
                        continue
                    c = fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
                    if region.contains(c) and r.size <= self.body * 1.1:
                        used.add(i)
                pix = page.get_pixmap(clip=region, dpi=170)
                items.append(Item("figure", pno, region, _side(region.x0, region.x1, self.gutter),
                                  block=Block("figure", image=pix.tobytes("png"),
                                              width_pt=region.width)))
            elif _TAB_CAP.match(sq):
                trows = self._table_rows(cap, rows, rules)
                if not trows:
                    continue
                blk = _table_block([rows[i] for i in trows], rules)
                if blk is None:
                    continue
                used.update(trows)
                rect = fitz.Rect(min(rows[i].x0 for i in trows), min(rows[i].y0 for i in trows),
                                 max(rows[i].x1 for i in trows), max(rows[i].y1 for i in trows))
                items.append(Item("table", pno, rect, _side(rect.x0, rect.x1, self.gutter), block=blk))
        # graphics nobody claimed: figures whose caption is part of the image
        taken = [it.rect for it in items]
        for f in figs:
            if f.width * f.height < 6000 or any(f.intersects(t) for t in taken):
                continue
            region = f & page.rect
            for i, r in enumerate(rows):
                c = fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
                if region.contains(c) and not self._is_caption(r):
                    used.add(i)
            pix = page.get_pixmap(clip=region, dpi=170)
            items.append(Item("figure", pno, region, _side(region.x0, region.x1, self.gutter),
                              block=Block("figure", image=pix.tobytes("png"),
                                          width_pt=region.width)))
            taken.append(region)

        # ruled grids with no caption: a table continued from the last page
        claimed = [it.rect for it in items if it.kind == "table"]
        vr = sorted((r for r in rules if r.width < 2.2), key=lambda r: r.y0)
        grids: list[fitz.Rect] = []
        for r in vr:
            if any(c.intersects(fitz.Rect(r.x0 - 3, r.y0, r.x1 + 3, r.y1)) for c in claimed):
                continue
            for g in grids:
                if r.y0 <= g.y1 + 3 and r.y1 >= g.y0 - 3:
                    g |= r
                    break
            else:
                grids.append(fitz.Rect(r))
        for g in grids:
            hs = [h for h in rules if h.width >= 2.2 and g.y0 - 3 <= h.y0 <= g.y1 + 3]
            if len(hs) < 2:
                continue
            g = fitz.Rect(min(h.x0 for h in hs), g.y0, max(h.x1 for h in hs), g.y1)
            trows = [i for i, r in enumerate(rows) if i not in used and
                     g.contains(fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2))]
            if not trows:
                continue
            blk = _table_block([rows[i] for i in trows], rules)
            if blk is None:
                continue
            used.update(trows)
            items.append(Item("table", pno, g, _side(g.x0, g.x1, self.gutter), block=blk))

        for i, r in enumerate(rows):
            if i in used:
                continue
            if self._is_caption(r):
                r.kind = "caption"
            rect = fitz.Rect(r.x0, r.y0, r.x1, r.y1)
            items.append(Item("row", pno, rect, _side(r.x0, r.x1, self.gutter), row=r))
        return self._order(items)

    def _is_caption(self, r: Row) -> bool:
        sq = r.squashed
        if not (_FIG_CAP.match(sq) or _TAB_CAP.match(sq) or _EN_CAP.match(r.text)):
            return False
        if len(sq) > 90:
            return False
        if r.size < self.body * 0.96:
            return True
        # body-size caption: centred within its column, not a sentence
        # that happens to start with "图 1 所示"
        side = _side(r.x0, r.x1, self.gutter)
        lo, hi = {"L": (self.left_l, self.gutter), "R": (self.gutter, self.right_r)}.get(
            side, (self.left_l, self.right_r))
        mid = (lo + hi) / 2
        return abs((r.x0 + r.x1) / 2 - mid) < 25 and r.x0 - lo > 2 * r.size \
            and not re.search(r"所示|可见|给出|中$", sq[:12])

    def _figure_region(self, cap: Row, figs, rows) -> fitz.Rect | None:
        cands = [f for f in figs if f.y1 <= cap.y0 + 6 and cap.y0 - f.y1 < 60
                 and f.x1 > cap.x0 - 80 and f.x0 < cap.x1 + 80]
        if not cands:
            return None
        region = fitz.Rect(max(cands, key=lambda f: f.y1))
        grew = True
        while grew:
            grew = False
            for f in figs:
                if f in cands or region.contains(f):
                    continue
                if f.x1 > region.x0 and f.x0 < region.x1 and \
                        f.y1 <= cap.y0 + 6 and region.y0 - f.y1 < 14:
                    region |= f
                    cands.append(f)
                    grew = True
        # sub-captions and labels between graphics and caption belong to it
        region.y1 = max(region.y1, cap.y0 - 1)
        return region & self.doc[cap.page].rect

    def _table_rows(self, cap: Row, rows, rules) -> list[int]:
        below = sorted((r for r in rules if r.y0 >= cap.y1 - 2 and r.x1 > cap.x0 and r.x0 < cap.x1),
                       key=lambda r: r.y0)
        bottom = None
        if below and below[0].y0 - cap.y1 < 40:
            top_rule = below[0]
            lo, hi = top_rule.x0 - 2, top_rule.x1 + 2
            bottom = top_rule.y0
            for r in below[1:]:
                if r.y0 - bottom > 600 or r.x1 < lo + 20 or r.x0 > hi - 20:
                    break
                if any(bottom < x.y0 < r.y0 and x.x0 >= lo - 2 and x.x1 <= hi + 2
                       and (self._is_caption(x) or
                            x.size >= self.body * 0.93 and len(x.squashed) > 12)
                       for x in rows):
                    break  # body text between the rules: the table ended
                bottom = r.y0
                lo, hi = min(lo, r.x0 - 2), max(hi, r.x1 + 2)
            if bottom - top_rule.y0 < 8:
                bottom = None  # a single rule: three-line table cut by a page
        if bottom is None:
            side = _side(cap.x0, cap.x1, self.gutter)
            lo, hi = {"L": (0, self.gutter), "R": (self.gutter, 9999)}.get(side, (0, 9999))
        idx = sorted((i for i, r in enumerate(rows)
                      if r.y0 > cap.y0 + 1 and r.x0 >= lo - 2 and r.x1 <= hi + 2),
                     key=lambda i: rows[i].y0)
        out = []
        for i in idx:
            r = rows[i]
            if _EN_CAP.match(r.text) and not out:
                continue  # the English caption line under the Chinese one
            if bottom is not None:
                if r.y0 > bottom + 2:
                    break
            elif r.size >= self.body * 0.93 and len(r.squashed) > 6:
                break
            out.append(i)
        return out

    def _order(self, items: list[Item]) -> list[Item]:
        full = sorted([it for it in items if it.side == "F"], key=lambda it: it.rect.y0)
        cols = [it for it in items if it.side != "F"]
        out: list[Item] = []
        prev_y = -1.0
        for f in full + [None]:
            lim = f.rect.y0 if f is not None else 1e9
            band = [it for it in cols if prev_y - 1 <= (it.rect.y0 + it.rect.y1) / 2 < lim]
            for side in ("L", "R"):
                out.extend(sorted((it for it in band if it.side == side),
                                  key=lambda it: it.rect.y0))
            for it in band:
                cols.remove(it)
            if f is not None:
                out.append(f)
                prev_y = (f.rect.y0 + f.rect.y1) / 2
        out.extend(sorted(cols, key=lambda it: it.rect.y0))
        return out

    # -- items -> blocks --------------------------------------------------
    def blocks(self) -> list[Block]:
        doc = self.doc
        # column edges from body rows
        lefts_l, lefts_r, rights_r = Counter(), Counter(), Counter()
        for pno in range(1, min(len(doc), 6)):
            for r in build_rows(page_chars(doc[pno]), self.gutter, self.body):
                if abs(r.size - self.body) > 0.3 or len(r.chars) < 10:
                    continue
                s = _side(r.x0, r.x1, self.gutter)
                if s == "L":
                    lefts_l[round(r.x0)] += 1
                elif s == "R":
                    lefts_r[round(r.x0)] += 1
                    rights_r[round(r.x1)] += 1
        self.left_l = lefts_l.most_common(1)[0][0]
        self.left_r = lefts_r.most_common(1)[0][0]
        self.right_r = rights_r.most_common(1)[0][0]
        self.right_l = self.left_l + (self.right_r - self.left_r)

        items: list[Item] = []
        for pno in range(len(doc)):
            items.extend(self.page_items(pno))

        blocks: list[Block] = []
        heading_nums: list[int] = []
        cur: Block | None = None
        last: Row | None = None
        mode = "front"
        title_seen = False
        ref_next = 1
        ref_size: float | None = None
        last_ref: Block | None = None
        skip_at: tuple[int, str] | None = None
        foot: tuple[int, str, float] | None = None
        pagefoot: tuple[int, str, float] | None = None
        page0_max = max((it.row.size for it in items
                         if it.kind == "row" and it.page == 0 and it.row.text.strip()), default=0)

        def flush():
            nonlocal cur
            if cur is not None and cur.text.strip():
                blocks.append(cur)
            cur = None

        def edges(r: Row) -> tuple[float, float]:
            s = _side(r.x0, r.x1, self.gutter)
            if s == "L":
                return self.left_l, self.right_l
            if s == "R":
                return self.left_r, self.right_r
            return self.left_l, self.right_r

        for it in items:
            if mode == "end":
                break
            if it.kind in ("figure", "table"):
                if mode == "refs":
                    continue
                flush()
                blocks.append(it.block)
                last = None
                continue
            r = it.row
            txt = r.text.strip()
            sq = r.squashed
            if not txt:
                continue

            # ---- reference list ------------------------------------------
            if mode in ("refs", "refskip"):
                m = re.match(r"^\[?(\d{1,3})\]?\s*\S", sq)
                if m and int(m.group(1)) == ref_next and \
                        (ref_size is None or abs(r.size - ref_size) < 0.4):
                    mode = "refs"
                    flush()
                    if ref_size is None:
                        ref_size = r.size
                    body = re.sub(r"^\[?\d{1,3}\]?\s*", "", unicodedata.normalize("NFKC", txt))
                    cur = last_ref = plain("ref", f"[{ref_next}] {body}")
                    ref_next += 1
                    continue
                if _BIO.search(txt) or r.size > self.body * 1.15 or \
                        re.match(r"^(Abstract|Key\s*words)", txt) or \
                        (ref_size is not None and abs(r.size - ref_size) > 0.4):
                    # author bios / English back matter: skip this column;
                    # the list may continue in the next one
                    mode = "refskip"
                    skip_at = (it.page, it.side)
                    continue
                if mode == "refskip" and (it.page, it.side) == skip_at:
                    continue
                mode = "refs"
                if last_ref is not None:
                    last_ref.runs = [(_join(last_ref.text, unicodedata.normalize("NFKC", txt)),
                                      False)]
                last = r
                continue

            if re.fullmatch(r"参考文献|References", sq):
                flush()
                blocks.append(plain("refhead", "参考文献"))
                mode = "refs"
                continue

            # ---- title -------------------------------------------------------
            if it.page == 0 and not title_seen:
                if r.size >= page0_max - 0.5 and any(is_cjk(c.c) for c in r.chars):
                    if cur is not None and cur.kind == "title":
                        cur.runs = [(cur.text + txt, False)]
                    else:
                        flush()
                        cur = plain("title", txt)
                    last = r
                    continue
                if cur is None or cur.kind != "title":
                    self.dropped.append(f"p0 masthead: {txt[:60]}")
                    continue
                title_seen = True

            # page-0 footnotes (收稿日期 / 基金项目 / 作者简介) run in a smaller
            # size at the foot of the left column: drop the whole block
            if it.page <= 1 and (_DROP_PARA.match(sq) or _BIO.search(txt)) \
                    and r.size < self.body * 0.95:
                # the paragraph above may run on into the next column
                foot = (it.page, it.side, r.y0)
                self.dropped.append(f"footnote: {txt[:60]}")
                continue
            if foot and it.page == foot[0] and r.y0 >= foot[2] - 1 \
                    and r.size < self.body * 0.95:
                self.dropped.append(f"footnote: {txt[:60]}")
                continue

            # page footnotes ("① http://…") at the foot of a column: a Word
            # manuscript keeps them as footnotes; they must not split the
            # paragraph that runs on into the next column
            H = self.doc[it.page].rect.height
            if re.match(r"^[\u2460-\u2473]", txt) and r.size < self.body * 0.9 \
                    and r.y0 > H * 0.5:
                pagefoot = (it.page, it.side, r.y0)
                self.dropped.append(f"page footnote: {txt[:60]}")
                continue
            if pagefoot and it.page == pagefoot[0] and it.side == pagefoot[1] \
                    and r.y0 >= pagefoot[2] and r.size < self.body * 0.9:
                continue

            # ---- captions ------------------------------------------------------
            if r.kind == "caption":
                if cur is not None and cur.kind == "caption" and last is not None \
                        and not (_EN_CAP.match(txt) and not _EN_CAP.match(cur.text)) \
                        and r.y0 - last.y1 < 1.2 * r.size and r.page == last.page:
                    cur.runs = [(_join(cur.text, txt), False)]
                else:
                    flush()
                    cur = plain("caption", txt)
                last = r
                continue
            if cur is not None and cur.kind == "caption" and last is not None \
                    and r.page == last.page and 0 <= r.y0 - last.y1 < 0.8 * r.size \
                    and abs(r.size - last.size) < 0.3 and r.size < self.body * 0.96:
                cur.runs = [(_join(cur.text, txt), False)]  # wrapped caption
                last = r
                continue

            # ---- headings ------------------------------------------------------
            left, right = edges(r)
            hm = _HEAD.match(re.sub(r"(?<=\d)．(?=\d)", ".", txt))
            short = r.x1 - r.x0 < 0.8 * (right - left) or r.size > self.body * 1.04
            if hm and short and not txt.endswith(_TERMINAL[:4]):
                nums = [int(x) for x in hm.group(1).split(".")]
                words = len(re.findall(r"[\u4e00-\u9fff]", hm.group(2))) >= 2 or \
                    len(re.findall(r"[A-Za-z]", hm.group(2))) >= 4
                if words and _plausible(heading_nums, nums):
                    flush()
                    heading_nums = nums
                    title = re.sub(r"(?<=[一-鿿])\s+(?=[一-鿿])", "",
                                   hm.group(2).strip())
                    blocks.append(plain("heading", f"{hm.group(1)} {title}", level=min(len(nums), 3)))
                    mode = "body"
                    last = None
                    continue
            if _UNNUMBERED.match(sq) and short:
                flush()
                blocks.append(plain("heading", sq, level=1))
                mode = "body"
                last = None
                continue

            if mode == "front" and it.side != "F" and abs(r.size - self.body) < 0.3 \
                    and len(sq) > 12:
                mode = "body"  # an introduction with no heading

            # ---- paragraphs ----------------------------------------------------
            if cur is not None and cur.kind in ("para", "front") and last is not None \
                    and r.page == last.page and abs(r.base - last.base) < 1.0 \
                    and r.x0 > last.x1:
                cur.runs = cur.runs + [(" ", False)] + list(r.runs())
                last = r
                continue
            new = cur is None or cur.kind not in ("para", "front")
            centred = r.x0 - left > 2.5 * r.size and right - r.x1 > 2.5 * r.size \
                and r.x1 - r.x0 < 0.7 * (right - left)
            if not new and last is not None:
                if _FORCE.match(sq):
                    new = True
                elif abs(r.size - last.size) > 0.8:
                    new = True
                elif it.side == "F" and r.page == last.page:
                    new = r.x0 > last.x0 + 1.2 * r.size or \
                        (centred or last.centred) or \
                        r.y0 - last.y1 > 1.2 * r.size
                elif r.x0 - left >= 1.2 * r.size and not (last.page != r.page or
                                                        _side(last.x0, last.x1, self.gutter) != it.side) \
                        and r.y0 - last.y1 > -2:
                    new = True
                elif r.x0 - left >= 1.2 * r.size and \
                        (last.page != r.page or _side(last.x0, last.x1, self.gutter) != it.side):
                    # indented first row of a column: new paragraph unless the
                    # previous one stopped mid-sentence
                    new = cur.text.rstrip().endswith(_TERMINAL)
                elif centred or last.centred:
                    new = True
                elif r.page == last.page and r.y0 - last.y1 > 1.2 * r.size:
                    new = True
            if new:
                flush()
                kind = "front" if mode == "front" else "para"
                cur = Block(kind, list(r.runs()), center=centred and mode == "front")
            else:
                cur.runs = _join_runs(cur.runs, r.runs())
            r.centred = centred
            last = r
        flush()

        blocks = _stitch(blocks)
        # drop footnotes / mastheads that survived as paragraphs; fold stray
        # punctuation (a full stop after a display formula) into its paragraph
        out = []
        for b in blocks:
            if b.kind == "para" and re.fullmatch(r"[\s.,;:()。，．；：（）]+", b.text) and out \
                    and out[-1].kind in ("para", "front"):
                out[-1].runs = _join_runs(out[-1].runs, b.runs)
                continue
            if b.kind in ("para", "front") and (_DROP_PARA.match(b.text.strip())
                                               or _BIO.search(b.text[:80])):
                self.dropped.append(f"para: {b.text[:60]}")
                continue
            if b.kind == "front":
                b.runs = [(t.strip("[]"), True) if sup else (t, sup) for t, sup in b.runs]
            out.append(b)
        return out


def _stitch(blocks: list[Block]) -> list[Block]:
    """Rejoin paragraphs that a column break, a page break or a float cut in
    two: a paragraph that stops mid-sentence takes the next paragraph in,
    and the figures / tables / captions in between move after it."""
    out: list[Block] = []
    pending: list[Block] = []
    open_para: Block | None = None

    def is_open(b: Block) -> bool:
        t = b.text.rstrip()
        return (b.kind == "para" and len(t) >= 30 and not b.center
                and not t.endswith(_TERMINAL) and not re.search(r"[(（]\d+[)）]$", t))

    for b in blocks:
        if open_para is not None and b.kind in ("figure", "table", "caption"):
            pending.append(b)
            continue
        if open_para is not None and b.kind == "para":
            starts_new = re.match(r"^\s*([(（]?\d{1,2}[)）]|[\u2460-\u2473]|\d{1,2}[)）]|"
                                  r"[a-z][)）]|图\s*\d|表\s*\d)", b.text)
            if not starts_new or re.match(r"^[．。，、；：）)]", b.text):
                open_para.runs = _join_runs(open_para.runs, b.runs)
                if not is_open(open_para):
                    out.extend(pending)
                    pending, open_para = [], None
                continue
        out.extend(pending)
        pending = []
        out.append(b)
        open_para = b if is_open(b) else None
    out.extend(pending)
    return out


# ---------------------------------------------------------------------------
# HTML engine (软件学报)
# ---------------------------------------------------------------------------

def _html_runs(el) -> list[tuple[str, bool]]:
    runs: list[tuple[str, bool]] = []

    def add(t: str, sup: bool):
        t = nk(t).replace("\xa0", " ")
        if not t:
            return
        if runs and runs[-1][1] == sup:
            runs[-1] = (runs[-1][0] + t, sup)
        else:
            runs.append((t, sup))

    def walk(e, sup=False):
        if e.tag == "sup":
            add(re.sub(r"\s+", "", e.text_content()), True)
            if e.tail:
                add(e.tail, sup)
            return
        if e.tag == "br":
            add(" ", sup)
        elif e.text:
            add(e.text, sup)
        for c in e:
            walk(c, sup)
        if e.tail:
            add(e.tail, sup)

    if el.text:
        add(el.text, False)
    for c in el:
        walk(c)
    return [(re.sub(r"\s+", " ", t), s) for t, s in runs]


def _fetch(url: str, dest: Path) -> bytes:
    if not dest.exists():
        dest.parent.mkdir(parents=True, exist_ok=True)
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        dest.write_bytes(urllib.request.urlopen(req, timeout=60).read())
    return dest.read_bytes()


def html_blocks(cfg: dict) -> list[Block]:
    base = cfg["url"].rsplit("/", 1)[0] + "/"
    raw = _fetch(cfg["url"], HERE / cfg["cache"])
    t = lxml.html.fromstring(raw.decode("utf-8-sig"))
    blocks: list[Block] = []
    titles = [e.text_content().strip() for e in t.xpath('//*[@class="article_title_cn"]')]
    blocks.append(plain("title", titles[0]))
    authors: list[tuple[str, bool]] = []
    for k, a in enumerate(t.xpath('//span[@class="author_CN"]')):
        if k:
            authors.append(("，", False))
        authors.append((a.text.strip(), False))
        sup = a.xpath("./sup")
        if sup:
            authors.append((sup[0].text_content().strip(), True))
    blocks.append(Block("front", authors, center=True))
    affs = [re.sub(r"\s+", " ", e.text_content()).strip()
            for e in t.xpath('//span[starts-with(@class,"danwei") and contains(@class,"_CN")]')]
    blocks.append(plain("front", "（" + "；".join(affs) + "）", center=True))
    abs_ = t.xpath('//div[@class="abs"]')
    keys = t.xpath('//div[@class="key"]')

    def key_text(k):
        words = [w.text_content().strip() for w in k.xpath('.//span[@class="keyword"]')]
        head = k.xpath('.//span[contains(@class,"key_title")]')[0].text_content().strip()
        return head, words

    blocks.append(plain("front", re.sub(r"\s+", " ", abs_[0].text_content()).strip()
                        .replace("摘要:", "摘要：", 1)))
    h, w = key_text(keys[0])
    blocks.append(plain("front", f"{h}：" + "；".join(w)))
    blocks.append(plain("front", titles[1], center=True))
    blocks.append(plain("front", re.sub(r"\s+", " ", abs_[1].text_content()).strip()))
    h, w = key_text(keys[1])
    blocks.append(plain("front", f"{h}: " + "; ".join(w)))

    body = t.xpath('//div[@class="article_body"]')[0]
    for e in body.iterchildren():
        cls = e.get("class") or ""
        if e.tag == "p":
            runs = _html_runs(e)
            if "".join(x for x, _ in runs).strip():
                runs[0] = (runs[0][0].lstrip(), runs[0][1])
                blocks.append(Block("para", runs))
        elif "paragraph_title" in cls:
            txt = re.sub(r"\s+", " ", e.text_content()).strip()
            blocks.append(plain("heading", txt, level=min(int(e.get("level", "1")), 3)))
        elif e.tag == "table":
            cells = [re.sub(r"\s+", " ", td.text_content()).strip() for td in e.xpath(".//td")]
            eq = cells[0].strip("$ ").strip()
            blocks.append(plain("formula", f"{eq}    {cells[1]}" if len(cells) > 1 else eq,
                                center=True))
        elif "figure" in cls:
            img = e.xpath(".//img")[0].get("src")
            data = _fetch(base + img, HERE / cfg["cache"].replace(".html", "") / Path(img).name)
            caps = [re.sub(r"\s+", " ", c.text_content()).strip()
                    for c in e.xpath('.//span[@class="caption"]') if c.text_content().strip()]
            blocks.append(Block("figure", image=data, width_pt=300))
            for c in sorted(dict.fromkeys(caps), key=lambda c: not c.startswith("图")):
                blocks.append(plain("caption", c))
        elif "table" in cls:
            caps = [re.sub(r"\s+", " ", c.text_content()).strip()
                    for c in e.xpath('.//span[@class="caption"]')]
            for c in sorted(dict.fromkeys(caps), key=lambda c: not c.startswith("表")):
                blocks.append(plain("caption", c))
            grid = []
            for tr in e.xpath('.//div[@class="table_content"]//table[1]//tr'):
                row = []
                for td in tr.xpath("./td|./th"):
                    row.append(re.sub(r"\s+", " ", nk(td.text_content())).strip())
                    row += [""] * (int(td.get("colspan", "1")) - 1)
                grid.append(row)
            width = max(len(r) for r in grid)
            blocks.append(Block("table", rows=[r + [""] * (width - len(r)) for r in grid]))

    blocks.append(plain("refhead", "参考文献"))
    rj = json.loads((HERE / cfg["refs_json"]).read_text())
    for i, txt in enumerate(rj["refs"], 1):
        blocks.append(plain("ref", f"[{i}] {txt}"))
    return blocks


# ---------------------------------------------------------------------------
# docx writer
# ---------------------------------------------------------------------------

def _font(style_or_run, latin: str, east: str, size: float | None = None, bold=None):
    f = style_or_run.font
    f.name = latin
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    el = style_or_run.element
    rpr = el.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for k in ("w:ascii", "w:hAnsi"):
        fonts.set(qn(k), latin)
    fonts.set(qn("w:eastAsia"), east)


def write_docx(blocks: list[Block], out: Path) -> dict:
    doc = Document()
    st = doc.styles
    _font(st["Normal"], "Times New Roman", "宋体", 10.5)
    st["Normal"].paragraph_format.space_after = Pt(0)
    st["Normal"].paragraph_format.line_spacing = 1.3
    _font(st["Title"], "Times New Roman", "黑体", 18, True)
    st["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st["Title"].paragraph_format.space_after = Pt(12)
    for lvl, size in ((1, 14), (2, 12), (3, 10.5)):
        s = st[f"Heading {lvl}"]
        _font(s, "Times New Roman", "黑体", size, True)
        s.font.italic = False
        s.paragraph_format.space_before = Pt(12 if lvl == 1 else 6)
        s.paragraph_format.space_after = Pt(6 if lvl == 1 else 3)
    _font(st["Caption"], "Times New Roman", "宋体", 9, False)
    st["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st["Caption"].paragraph_format.space_after = Pt(6)

    stats = Counter()
    for b in blocks:
        stats[b.kind] += 1
        if b.kind == "title":
            doc.add_paragraph(b.text, style="Title")
        elif b.kind == "heading":
            doc.add_paragraph(b.text, style=f"Heading {b.level}")
        elif b.kind == "refhead":
            doc.add_paragraph(b.text, style="Heading 1")
        elif b.kind == "caption":
            doc.add_paragraph(b.text, style="Caption")
        elif b.kind == "figure":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            width = min(b.width_pt, 430)
            p.add_run().add_picture(io.BytesIO(b.image), width=Pt(width))
        elif b.kind == "table":
            rows = b.rows or []
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, row in enumerate(rows):
                for j, txt in enumerate(row):
                    cell = t.cell(i, j)
                    cell.text = txt
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.font.bold = i == 0
            doc.add_paragraph()
        elif b.kind == "ref":
            p = doc.add_paragraph(b.text)
            p.paragraph_format.left_indent = Pt(21)
            p.paragraph_format.first_line_indent = Pt(-21)
            for run in p.runs:
                run.font.size = Pt(9)
        else:
            p = doc.add_paragraph()
            if b.center or b.kind == "formula":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif b.kind == "para":
                p.paragraph_format.first_line_indent = Pt(21)
            for txt, sup in b.runs:
                run = p.add_run(txt)
                if sup:
                    run.font.superscript = True
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return dict(stats)


# ---------------------------------------------------------------------------

JOBS = {
    "kg": dict(src="zh_papers/liuqiao_kg/manuscript.pdf", dir="liuqiao_kg",
               authors="刘峤  李杨  段宏  刘瑶  秦志光"),
    "odt": dict(src="zh_papers/luohuilan_odt/manuscript.pdf", dir="luohuilan_odt"),
    "tod": dict(src="zh_papers/zhaoyangyang_tod/manuscript.pdf", dir="zhaoyangyang_tod",
                authors="赵阳洋  王振宇  王佩  杨添  张睿  尹凯"),
    "yin": dict(src="zh_pdfs/aas_yinhongpeng_2016.pdf", dir="yinhongpeng_odt",
                refs_json="zh_refs/aas_yinhongpeng_odt.json",
                authors="尹宏鹏^1,2  陈波^2  柴毅^2  刘兆栋^2",
                affil="（1. 信息物理社会可信服务计算教育部重点实验室（重庆大学） 重庆 400030；"
                      "2. 重庆大学自动化学院 重庆 400044）"),
    "mrc": dict(html=True, dir="guyingjie_mrc",
                url="https://www.jos.org.cn/html/2020/7/6048.htm",
                cache="zh_html/jos_mrc_6048.html",
                refs_json="zh_refs/jos_mrc_6048.json"),
}


def _place_floats(blocks: list[Block]) -> list[Block]:
    """A table or figure goes after the paragraph that first mentions it, as
    an author places it in Word; the journal's page layout may print it
    earlier (top of the page). Floats are only moved later, never earlier."""
    units: list[tuple[int, int]] = []  # [start, end) runs of float blocks
    i = 0
    while i < len(blocks):
        if blocks[i].kind in ("figure", "table", "caption"):
            j = i
            while j < len(blocks) and blocks[j].kind in ("figure", "table", "caption"):
                j += 1
            units.append((i, j))
            i = j
        else:
            i += 1
    moves: dict[int, list[Block]] = {}
    drop: set[int] = set()
    for a, b in units:
        labels = [m.groups() for blk in blocks[a:b] if blk.kind == "caption"
                  if (m := re.match(r"^\s*([图表])\s*(\d+)", blk.text))]
        if not labels or not any(blk.kind != "caption" for blk in blocks[a:b]):
            continue
        pats = [re.compile(rf"{k}\s*{n}(?![\d.．])") for k, n in labels]
        first = next((k for k, blk in enumerate(blocks)
                      if blk.kind == "para" and any(p.search(blk.text) for p in pats)), None)
        if first is None or first < a:
            continue
        moves.setdefault(first, []).extend(blocks[a:b])
        drop.update(range(a, b))
    out: list[Block] = []
    for k, blk in enumerate(blocks):
        if k not in drop:
            out.append(blk)
        out.extend(moves.get(k, []))
    return out


def _set_authors(blocks: list[Block], authors: str, affil: str | None) -> list[Block]:
    """Replace the author line(s) under the title. Typeset author lines
    letter-space two-character names, so the names are given in JOBS;
    "^1,2" marks an affiliation superscript."""
    t = next(i for i, b in enumerate(blocks) if b.kind == "title")
    j = t + 1
    while j < len(blocks) and blocks[j].kind == "front" and not re.search(
            r"^[(（]|大学|学院|研究所|摘\s*要", blocks[j].text):
        j += 1
    runs: list[tuple[str, bool]] = []
    for part in re.split(r"(\^[\d,]+)", authors):
        if part.startswith("^"):
            runs.append((part[1:], True))
        elif part:
            runs.append((part, False))
    new = [Block("front", runs, center=True)]
    if affil:
        new.append(plain("front", affil, center=True))
    return blocks[:t + 1] + new + blocks[j:]


def build(key: str, verbose: bool = False) -> None:
    cfg = JOBS[key]
    out = HERE / "zh_papers" / cfg["dir"] / "manuscript.docx"
    if cfg.get("html"):
        blocks = html_blocks(cfg)
        dropped: list[str] = []
    else:
        lay = PdfLayout(HERE / cfg["src"], cfg)
        blocks = lay.blocks()
        dropped = lay.dropped
        if cfg.get("refs_json"):
            d = json.loads((HERE / cfg["refs_json"]).read_text())
            blocks = [b for b in blocks if b.kind != "ref"]
            blocks += [plain("ref", f"[{r['n']}] {r['text']}") for r in d["refs"]]
    if cfg.get("authors"):
        blocks = _set_authors(blocks, cfg["authors"], cfg.get("affil"))
    if not cfg.get("html"):
        blocks = _place_floats(blocks)
    stats = write_docx(blocks, out)
    heads = [b.text for b in blocks if b.kind == "heading"]
    print(f"{cfg['dir']}: {stats}")
    print("  headings:", heads)
    if verbose:
        for d in dropped:
            print("  dropped:", d)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    for k in args or list(JOBS):
        build(k, verbose="-v" in sys.argv)
